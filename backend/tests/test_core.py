import unittest
from io import BytesIO

from docx import Document as DocxDocument

from app.core import DocumentError, answer_question, make_summary, parse_document, quick_revision


SAMPLE = b"""Chapter 1: Neural Networks
Neural networks learn patterns from examples. A training dataset contains examples and labels.
The model adjusts its weights during training to reduce prediction error.

Chapter 2: Evaluation
Precision measures the fraction of predicted positives that are correct. Recall measures the fraction of actual positives that are found.
"""


class DocumentCoreTests(unittest.TestCase):
    def setUp(self):
        self.document = parse_document("notes.txt", SAMPLE)

    def test_upload_rejects_unsupported_and_binary_content(self):
        with self.assertRaisesRegex(DocumentError, "Unsupported file type"):
            parse_document("notes.exe", b"not a document")
        with self.assertRaisesRegex(DocumentError, "binary data"):
            parse_document("notes.txt", b"bad\x00text")

    def test_sections_and_chunks_keep_source_locations(self):
        self.assertEqual([section.title for section in self.document.sections], ["Chapter 1: Neural Networks", "Chapter 2: Evaluation"])
        self.assertEqual(self.document.chunks[0].section, "Chapter 1: Neural Networks")
        self.assertIsNone(self.document.chunks[0].page)

    def test_docx_headings_and_paragraphs_are_extracted(self):
        source = DocxDocument()
        source.add_heading("Chapter 1: Neural Networks", level=1)
        source.add_paragraph("Neural networks learn patterns from examples and adjust their weights.")
        stream = BytesIO()
        source.save(stream)

        parsed = parse_document("notes.docx", stream.getvalue())
        self.assertEqual(parsed.sections[0].title, "Chapter 1: Neural Networks")
        self.assertIn("learn patterns", parsed.chunks[0].text)

    def test_question_answers_quote_document_and_cite_source(self):
        result = answer_question(self.document, "How does a neural network learn patterns?")
        self.assertTrue(result["grounded"])
        self.assertIn("learn patterns", result["answer"])
        self.assertEqual(result["citations"][0]["document"], "notes.txt")
        self.assertEqual(result["citations"][0]["section"], "Chapter 1: Neural Networks")

    def test_unsupported_question_is_refused(self):
        result = answer_question(self.document, "What is the capital of France?")
        self.assertFalse(result["grounded"])
        self.assertEqual(result["citations"], [])

    def test_exam_summary_and_time_revision_are_source_grounded(self):
        summary = make_summary(self.document, intent="university_exam", learning_level="beginner")
        self.assertTrue(summary["exam_material"]["notice"].startswith("Potential questions"))
        self.assertTrue(all(point["source"]["document"] == "notes.txt" for point in summary["key_points"]))
        revision = quick_revision(self.document, 10)
        self.assertEqual(revision["minutes"], 10)
        self.assertTrue(revision["priority_points"])
        with self.assertRaisesRegex(DocumentError, "duration"):
            quick_revision(self.document, 7)


if __name__ == "__main__":
    unittest.main()