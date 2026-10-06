from __future__ import annotations

import io
import re
import zipfile
from collections import Counter
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

from docx import Document as DocxDocument
from pypdf import PdfReader


MAX_UPLOAD_BYTES = 20 * 1024 * 1024
SUPPORTED_TYPES = {".pdf", ".docx", ".txt"}
STOP_WORDS = set(
    "a an and are as at be been being by can could for from has have in is it its "
    "of on or that the their this to was were will with which when where who why "
    "what how into than then them these those through using used use also such "
    "may might should would not but if we you they he she his her our your".split()
)
WORD_RE = re.compile(r"[A-Za-z][A-Za-z0-9'-]{2,}")
SENTENCE_RE = re.compile(r"(?<=[.!?])\s+|\n+")


class DocumentError(ValueError):
    pass


@dataclass(frozen=True)
class Section:
    title: str
    text: str
    page: int | None = None


@dataclass(frozen=True)
class SourceChunk:
    id: str
    text: str
    section: str
    page: int | None

    def citation(self, document_name: str) -> dict[str, Any]:
        return {
            "document": document_name,
            "page": self.page,
            "section": self.section,
            "chunk_id": self.id,
        }


@dataclass(frozen=True)
class ParsedDocument:
    name: str
    file_type: str
    sections: list[Section]
    chunks: list[SourceChunk]


def parse_document(name: str, payload: bytes) -> ParsedDocument:
    suffix = Path(name).suffix.lower()
    if suffix not in SUPPORTED_TYPES:
        raise DocumentError("Unsupported file type. Upload a PDF, DOCX, or TXT file.")
    if not payload:
        raise DocumentError("The uploaded file is empty.")
    if len(payload) > MAX_UPLOAD_BYTES:
        raise DocumentError("The file exceeds the 20 MB upload limit.")

    if suffix == ".pdf":
        pages = _read_pdf(payload)
    elif suffix == ".docx":
        pages = _read_docx(payload)
    else:
        pages = _read_text(payload)

    sections = _make_sections(pages)
    if not sections or not any(section.text.strip() for section in sections):
        raise DocumentError(
            "No readable text was found. Scanned PDFs need OCR before they can be processed."
        )
    chunks = _make_chunks(sections)
    return ParsedDocument(name, suffix[1:], sections, chunks)


def _read_pdf(payload: bytes) -> list[tuple[int, str]]:
    if not payload.startswith(b"%PDF-"):
        raise DocumentError("The file extension and PDF contents do not match.")
    try:
        reader = PdfReader(io.BytesIO(payload), strict=True)
        if reader.is_encrypted:
            raise DocumentError("Password-protected PDFs are not supported.")
        return [(index + 1, page.extract_text() or "") for index, page in enumerate(reader.pages)]
    except DocumentError:
        raise
    except Exception as exc:
        raise DocumentError("The PDF could not be read. It may be corrupted.") from exc


def _read_docx(payload: bytes) -> list[tuple[int | None, str]]:
    try:
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            entries = archive.infolist()
            if "word/document.xml" not in archive.namelist():
                raise DocumentError("The file extension and DOCX contents do not match.")
            if len(entries) > 10_000 or sum(entry.file_size for entry in entries) > 100 * 1024 * 1024:
                raise DocumentError("The DOCX expands beyond the safe processing limit.")
        document = DocxDocument(io.BytesIO(payload))
    except DocumentError:
        raise
    except Exception as exc:
        raise DocumentError("The DOCX could not be read. It may be corrupted.") from exc
    paragraphs = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            heading = paragraph.style is not None and paragraph.style.name.startswith("Heading")
            paragraphs.append(("# " if heading else "") + text)
    return [(None, "\n".join(paragraphs))]


def _read_text(payload: bytes) -> list[tuple[int | None, str]]:
    if b"\x00" in payload:
        raise DocumentError("The TXT file contains binary data and cannot be processed.")
    try:
        text = payload.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise DocumentError("TXT files must use UTF-8 encoding.") from exc
    return [(None, text)]


def _looks_like_heading(line: str) -> bool:
    cleaned = line.strip().lstrip("# ").strip()
    if not cleaned or len(cleaned) > 100 or len(cleaned.split()) > 12:
        return False
    return line.lstrip().startswith("#") or bool(
        re.match(r"^(chapter|section|unit|lecture)\s+[\w.-]+", cleaned, re.I)
    ) or (cleaned.isupper() and len(cleaned.split()) <= 10)


def _make_sections(pages: list[tuple[int | None, str]]) -> list[Section]:
    result: list[Section] = []
    current_title = "Document"
    current_page: int | None = None
    current_lines: list[str] = []

    def flush() -> None:
        text = "\n".join(current_lines).strip()
        if text:
            result.append(Section(current_title, text, current_page))

    for page, text in pages:
        if not text.strip():
            continue
        if page is not None:
            current_page = page
        for line in text.splitlines():
            line = line.strip()
            if not line:
                continue
            if _looks_like_heading(line):
                flush()
                current_title = line.lstrip("# ").strip()
                current_lines = []
                if page is not None:
                    current_page = page
            else:
                current_lines.append(line)
        flush()
        current_lines = []
    return result


def _make_chunks(sections: list[Section], chunk_words: int = 480, overlap: int = 60) -> list[SourceChunk]:
    chunks: list[SourceChunk] = []
    for section in sections:
        words = section.text.split()
        start = 0
        while start < len(words):
            end = min(start + chunk_words, len(words))
            chunks.append(
                SourceChunk(
                    id=f"chunk-{len(chunks) + 1}",
                    text=" ".join(words[start:end]),
                    section=section.title,
                    page=section.page,
                )
            )
            if end == len(words):
                break
            start = end - overlap
    return chunks


def _tokens(text: str) -> list[str]:
    return [word.lower() for word in WORD_RE.findall(text) if word.lower() not in STOP_WORDS]


def _sentences(text: str) -> list[str]:
    return [sentence.strip() for sentence in SENTENCE_RE.split(text) if len(_tokens(sentence)) >= 5]


def _rank_sentences(chunks: list[SourceChunk]) -> list[tuple[float, str, SourceChunk]]:
    candidates = [(sentence, chunk) for chunk in chunks for sentence in _sentences(chunk.text)]
    frequency = Counter(token for sentence, _ in candidates for token in _tokens(sentence))
    maximum = max(frequency.values(), default=1)
    ranked = []
    seen = set()
    for sentence, chunk in candidates:
        normalized = " ".join(_tokens(sentence))
        if normalized in seen:
            continue
        seen.add(normalized)
        terms = _tokens(sentence)
        length_penalty = 1 + abs(len(terms) - 24) / 80
        score = sum(frequency[token] / maximum for token in terms) / length_penalty
        ranked.append((score, sentence, chunk))
    return sorted(ranked, key=lambda item: item[0], reverse=True)


def _topic(sentence: str) -> str:
    terms = _tokens(sentence)
    return max(terms, key=lambda term: (len(term), -terms.index(term))) if terms else "General"


def make_summary(
    document: ParsedDocument,
    mode: str = "standard",
    learning_level: str = "intermediate",
    intent: str = "general",
) -> dict[str, Any]:
    modes = {"concise": 4, "standard": 8, "detailed": 14, "key_points": 8}
    if mode not in modes:
        raise DocumentError("Summary mode must be concise, standard, detailed, or key_points.")
    if learning_level not in {"beginner", "intermediate", "advanced"}:
        raise DocumentError("Learning level must be beginner, intermediate, or advanced.")
    if intent not in {"general", "university_exam", "viva", "interview", "assignment"}:
        raise DocumentError("Unsupported learning intent.")

    ranked = _rank_sentences(document.chunks)
    selected = ranked[: modes[mode]]
    points = [
        {"text": sentence, "topic": _topic(sentence), "source": chunk.citation(document.name)}
        for _, sentence, chunk in selected
    ]
    output: dict[str, Any] = {
        "document": document.name,
        "mode": mode,
        "learning_level": learning_level,
        "intent": intent,
        "key_points": points,
        "sections": [
            {
                "title": section.title,
                "page": section.page,
                "summary": " ".join(
                    sentence
                    for _, sentence, chunk in ranked
                    if chunk.section == section.title and chunk.page == section.page
                )[:1200],
            }
            for section in document.sections
        ],
        "source_coverage": "Every listed point is extracted verbatim from the uploaded document.",
    }
    if intent == "university_exam":
        output["exam_material"] = {
            "definitions": [point["text"] for point in points[:4]],
            "potential_questions": [f"Explain: {point['topic']}" for point in points[:5]],
            "notice": "Potential questions based on the uploaded material; not predictions of an actual exam.",
        }
    elif intent == "viva":
        output["viva_questions"] = [f"What is {point['topic']}?" for point in points[:6]]
    elif intent == "interview":
        output["interview_questions"] = [f"How is {point['topic']} used?" for point in points[:6]]
    elif intent == "assignment":
        output["assignment_notes"] = [point["text"] for point in points]
    return output


def answer_question(document: ParsedDocument, question: str) -> dict[str, Any]:
    query_terms = set(_tokens(question))
    if not query_terms:
        raise DocumentError("Enter a question about the document.")
    ranked = []
    for chunk in document.chunks:
        terms = set(_tokens(chunk.text))
        overlap = len(query_terms & terms)
        score = overlap / len(query_terms)
        if overlap:
            ranked.append((score, overlap, chunk))
    ranked.sort(key=lambda item: (item[0], item[1]), reverse=True)
    if not ranked or ranked[0][0] < 0.3:
        return {
            "answer": "The uploaded document does not contain enough information to answer this question.",
            "grounded": False,
            "citations": [],
        }
    sources = [
        {"excerpt": chunk.text[:900], "citation": chunk.citation(document.name)}
        for _, _, chunk in ranked[:3]
    ]
    return {
        "answer": sources[0]["excerpt"],
        "grounded": True,
        "citations": [source["citation"] for source in sources],
        "evidence": sources,
    }


def quick_revision(document: ParsedDocument, minutes: int) -> dict[str, Any]:
    if minutes not in {5, 10, 15, 30}:
        raise DocumentError("Quick revision duration must be 5, 10, 15, or 30 minutes.")
    ranked = _rank_sentences(document.chunks)
    counts = {5: 5, 10: 10, 15: 15, 30: 25}
    points = [
        {"text": sentence, "topic": _topic(sentence), "source": chunk.citation(document.name)}
        for _, sentence, chunk in ranked[: counts[minutes]]
    ]
    return {
        "document": document.name,
        "minutes": minutes,
        "priority_points": points,
        "self_test": [f"Explain {point['topic']} without looking at the notes." for point in points[:5]],
    }


def generate_quiz(document: ParsedDocument, count: int = 5) -> list[dict[str, Any]]:
    if not 1 <= count <= 20:
        raise DocumentError("Quiz question count must be between 1 and 20.")
    ranked = _rank_sentences(document.chunks)
    questions = []
    for _, sentence, chunk in ranked:
        terms = _tokens(sentence)
        if len(terms) < 6:
            continue
        frequent = Counter(terms)
        answer = max(terms, key=lambda term: (frequent[term], len(term)))
        if len(answer) < 5:
            continue
        blank = re.compile(rf"\b{re.escape(answer)}\b", re.I)
        prompt = blank.sub("______", sentence, count=1)
        questions.append(
            {
                "id": f"q-{len(questions) + 1}",
                "type": "short_answer",
                "prompt": prompt,
                "topic": answer,
                "source": chunk.citation(document.name),
                "_answer": answer.lower(),
            }
        )
        if len(questions) == count:
            break
    if not questions:
        raise DocumentError("There is not enough readable source material to create a quiz.")
    return questions


def score_quiz(questions: list[dict[str, Any]], answers: dict[str, str]) -> dict[str, Any]:
    results = []
    for question in questions:
        expected = question["_answer"]
        supplied = (answers.get(question["id"]) or "").strip().lower()
        correct = supplied == expected
        results.append(
            {
                "question_id": question["id"],
                "topic": question["topic"],
                "correct": correct,
                "expected_answer": expected,
                "source": question["source"],
            }
        )
    total = len(results)
    score = round(sum(result["correct"] for result in results) / total * 100) if total else 0
    return {"score_percent": score, "correct": sum(r["correct"] for r in results), "total": total, "results": results}


def public_question(question: dict[str, Any]) -> dict[str, Any]:
    return {key: value for key, value in question.items() if not key.startswith("_")}


def serialize_document(document: ParsedDocument) -> dict[str, Any]:
    return {
        "name": document.name,
        "file_type": document.file_type,
        "sections": [asdict(section) for section in document.sections],
        "chunks": [asdict(chunk) for chunk in document.chunks],
    }